#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""t6 渲染 v2（审计账本式纸墨系统）交付门禁：G1 文字 / G2 数字 / G3 设计 / G4 PPT / G5 PDF。
输出 gates_report_v2.json。"""
import hashlib, json, re, subprocess
from pathlib import Path

HERE = Path(__file__).resolve().parent
MD = HERE / "AI无佣与人机协同_正式稿_v1.md"
HTML = HERE / "AI无佣与人机协同_正式稿_v2.html"
PDF = HERE / "AI无佣与人机协同_正式稿_v2.pdf"
PPTX = HERE / "AI无佣与人机协同_汇报版_v2.pptx"
DETECT = Path("/root/zhijian/dsh-expert-library/knowledge/skills/finesse-ui/scripts/detect.mjs")
NUM = re.compile(r"\d+(?:\.\d+)?")
STRUCTURAL = {f"{i:02d}" for i in range(0, 14)} | {"98"}
FORBIDDEN = ["渠道灯塔", "统筹会长", "制度君", "海外瞭望员", "居住服务·Q", "居住服务·C",
             "居住服务·Z", "居住服务·H", "BK-019", "讨论稿", "v5.4", "修订清单", "评审轮次",
             "250 万亿", "核验约 40 项"]


def sha(p): return hashlib.sha256(Path(p).read_bytes()).hexdigest()


def visible(html):
    b = re.sub(r"(?s)<script.*?</script>|<style.*?</style>|<!--.*?-->", " ", html)
    return re.sub(r"(?s)<[^>]+>", " ", b)


md = MD.read_text(); html = HTML.read_text(); body = visible(html)
rep = {"gates": {}, "artifacts": {}, "design": {}}

# G1
hits = [t for t in FORBIDDEN if t in body]
rep["gates"]["G1_text"] = {"forbidden_hits": hits, "pass": not hits}

# G2
extra = sorted(set(NUM.findall(body)) - set(NUM.findall(md)) - STRUCTURAL, key=lambda x: (len(x), x))
from pptx import Presentation
prs = Presentation(PPTX)
ptxt = []
for sl in prs.slides:
    for sh in sl.shapes:
        if sh.has_text_frame: ptxt.append(sh.text_frame.text)
        if sh.has_table:
            for r in sh.table.rows:
                for c in r.cells: ptxt.append(c.text)
    if sl.has_notes_slide: ptxt.append(sl.notes_slide.notes_text_frame.text)
pt = "\n".join(ptxt)
extra_p = sorted(set(NUM.findall(pt)) - set(NUM.findall(md)) - STRUCTURAL, key=lambda x: (len(x), x))
rep["gates"]["G2_numbers"] = {"new_in_html": extra, "new_in_pptx": extra_p,
                              "pass": not extra and not extra_p}

# G3 设计（detector + 九档视口 + 图表几何 + 必需图卡齐备）
det = json.loads(subprocess.run(["node", str(DETECT), "--json", str(HTML)],
                                capture_output=True, text=True).stdout)
findings = [f for x in det.get("files", []) for f in x.get("findings", [])]
vp = json.loads((HERE / "validation_viewports_v2.json").read_text())
vp_fail = [w for w, v in vp.items() if not v.get("pass")]
def count(pat): return len(re.findall(pat, html))
tbodies = re.findall(r"(?s)<tbody>(.*?)</tbody>", html)
def rows(i): return len(re.findall(r"<tr>", tbodies[i])) if i < len(tbodies) else 0
exhibits = {
    "成本树八段表行数": rows(0),
    "成本树环节数": len(re.findall(r'data-label="环节"', tbodies[0] if tbodies else "")),
    "观察面板行数": rows(1),
    "责任矩阵格数": count(r'<div class="mx(?:"| )'),
    "责任矩阵护城河格数": count(r'<div class="mx moat">'),
    "三类场景卡数": count(r'<div class="sc">'),
    "三重剥离条数": count(r'<div class="strip(?: result)?">'),
    "编号图卡数": count(r'<span class="ex-no">'),
    "柱状图轴标签数": count(r'<div class="axlab">'),
    "三条洞察行数": count(r'<div class="keyrow">'),
}
anchors_ok = exhibits["编号图卡数"] == 3 and exhibits["柱状图轴标签数"] == 2
rep["gates"]["G3_design"] = {"detector_p0": det.get("p0"), "detector_findings": findings,
                             "viewports_failing": vp_fail, "viewports_tested": len(vp),
                             "axis_labels": exhibits["柱状图轴标签数"],
                             "pass": det.get("p0") == 0 and not findings and not vp_fail and anchors_ok}
rep["design"] = {"exhibits": exhibits, "exhibit_gate_pass": anchors_ok,
                 "ledger_tokens": {"page": "#F1EEE7", "card": "#FDFCF9",
                                   "accent": "#10605A", "gold": "#9A6B22"}}

# G4 PPT
pj = json.loads((HERE / "validation_pptx_structure_v2.json").read_text())
s = pj.get("summary", {})
rep["gates"]["G4_ppt"] = {"summary": s, "slides": len(prs.slides._sldIdLst),
                          "notes_on_all": all(x.has_notes_slide and x.notes_slide.notes_text_frame.text.strip()
                                              for x in prs.slides),
                          "pass": s.get("error", 1) == 0 and s.get("warning", 1) == 0}

# G5 PDF + 逐页栅格化留证
import fitz
doc = fitz.open(PDF)
mm = 2.834645
box = fitz.Rect(14 * mm, 15 * mm, (210 - 14) * mm, (297 - 15) * mm)
viol = 0
for pg in doc:
    for b in pg.get_text("blocks"):
        r = fitz.Rect(b[:4])
        if r.y0 > 805: continue
        if r.x0 < box.x0 - 3 or r.x1 > box.x1 + 3 or r.y0 < box.y0 - 3 or r.y1 > box.y1 + 3: viol += 1
FOOT = "98wiki ｜ 智见 / 行业研究报告"
n = doc.page_count
pngs = sorted((HERE / "shots_v2").glob("pdf_v2_p*.png"))
rep["gates"]["G5_pdf"] = {"pages": n, "content_box_violations": viol,
                         "footer_on_body_pages": all(FOOT in doc[i].get_text() for i in range(1, n - 1)),
                         "cover_unfooted": FOOT not in doc[0].get_text(),
                         "back_unfooted": FOOT not in doc[n - 1].get_text(),
                         "rasterized_pages": len(pngs),
                         "pass": viol == 0 and len(pngs) == n
                                 and all(FOOT in doc[i].get_text() for i in range(1, n - 1))}

for f in [MD.name, HTML.name, PDF.name, PPTX.name, "build_deck_v2.py", "gates_v2.py"]:
    p = HERE / f
    rep["artifacts"][f] = {"sha256": sha(p), "bytes": p.stat().st_size}
rep["all_pass"] = all(g["pass"] for g in rep["gates"].values())
(HERE / "gates_report_v2.json").write_text(json.dumps(rep, ensure_ascii=False, indent=1))
print(json.dumps({k: v["pass"] for k, v in rep["gates"].items()}, ensure_ascii=False))
print("new_in_html:", extra, "| new_in_pptx:", extra_p)
print("viewports:", len(vp), "| exhibits:", exhibits)
print("pdf pages:", n, "| rasterized:", len(pngs), "| all_pass:", rep["all_pass"])
