#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""t6 渲染 v2（审计账本式纸墨系统）交付门禁：G1 文字 / G2 数字 / G3 设计 / G4 PPT / G5 PDF。
输出 gates_report_v3.json。"""
import hashlib, json, re, subprocess
from pathlib import Path

HERE = Path(__file__).resolve().parent
MD = HERE / "AI无佣与人机协同_正式稿_v3.md"
HTML = HERE / "AI无佣与人机协同_正式稿_v3.html"
PDF = HERE / "AI无佣与人机协同_正式稿_v3.pdf"
PPTX = HERE / "AI无佣与人机协同_汇报版_v3.pptx"
DETECT = Path("/root/zhijian/dsh-expert-library/knowledge/skills/finesse-ui/scripts/detect.mjs")
NUM = re.compile(r"\d+(?:\.\d+)?")
STRUCTURAL = ({f"{i:02d}" for i in range(0, 20)} | {f"{i}" for i in range(1, 20)}
              | {"98", "2.1", "3.1", "3.2", "3.3", "3.4", "3.5", "3.6", "3.7", "3.8", "3.9", "4.1", "4.2", "4.3"})
# 注意：「250 万亿元」「核验约 40 项」按渲染要求出现在「未采用数字留痕」小节中，
# 那是防编造声明本身，不在禁例之列；禁例只覆盖过程痕迹与专家实名。
FORBIDDEN = ["渠道灯塔", "统筹会长", "制度君", "海外瞭望员", "居住服务·Q", "居住服务·C",
             "居住服务·Z", "居住服务·H", "BK-019", "讨论稿", "v5.4", "修订清单", "评审轮次"]


def sha(p): return hashlib.sha256(Path(p).read_bytes()).hexdigest()


def visible(html):
    b = re.sub(r"(?s)<script.*?</script>|<style.*?</style>|<!--.*?-->", " ", html)
    return re.sub(r"(?s)<[^>]+>", " ", b)


md = MD.read_text(); html = HTML.read_text(); body = visible(html)
rep = {"gates": {}, "artifacts": {}, "design": {}}

# G1
hits = [t for t in FORBIDDEN if t in body]
rep["gates"]["G1_text"] = {"forbidden_hits": hits, "pass": not hits}

# G2
extra = sorted(set(NUM.findall(body)) - set(NUM.findall(md)) - STRUCTURAL, key=lambda x: (len(x), x))
from pptx import Presentation
prs = Presentation(PPTX)
ptxt = []
for sl in prs.slides:
    for sh in sl.shapes:
        if sh.has_text_frame: ptxt.append(sh.text_frame.text)
        if sh.has_table:
            for r in sh.table.rows:
                for c in r.cells: ptxt.append(c.text)
    if sl.has_notes_slide: ptxt.append(sl.notes_slide.notes_text_frame.text)
pt = "\n".join(ptxt)
extra_p = sorted(set(NUM.findall(pt)) - set(NUM.findall(md)) - STRUCTURAL, key=lambda x: (len(x), x))
rep["gates"]["G2_numbers"] = {"new_in_html": extra, "new_in_pptx": extra_p,
                              "pass": not extra and not extra_p}

# G3 设计（detector + 九档视口 + 图表几何 + 必需图卡齐备）
det = json.loads(subprocess.run(["node", str(DETECT), "--json", str(HTML)],
                                capture_output=True, text=True).stdout)
findings = [f for x in det.get("files", []) for f in x.get("findings", [])]
vp = json.loads((HERE / "validation_viewports_v3.json").read_text())
vp_fail = [w for w, v in vp.items() if not v.get("pass")]
def count(pat): return len(re.findall(pat, html))
tbodies = re.findall(r"(?s)<tbody>(.*?)</tbody>", html)
def rows(i): return len(re.findall(r"<tr>", tbodies[i])) if i < len(tbodies) else 0
exhibits = {
    "表1 八段成本树行数": rows(0),
    "表2 责任边界矩阵行数": rows(1),
    "表3 三类场景行数": rows(2),
    "表4 Q1-Q5 预言表行数": rows(3),
    "表5 观察面板行数": rows(4),
    "责任链条断点数": count(r'<div class="nrow">'),
    "三重剥离条数": count(r'<div class="strip(?: result)?">'),
    "编号图卡数": count(r'<span class="ex-no">'),
    "柱状图轴标签数": count(r'<div class="axlab">'),
    "未采用数字留痕条目": count(r'①「存量住房资产'),
}
anchors_ok = (exhibits["表1 八段成本树行数"] == 8 and exhibits["表2 责任边界矩阵行数"] == 3
              and exhibits["表3 三类场景行数"] == 3 and exhibits["表4 Q1-Q5 预言表行数"] == 5
              and exhibits["表5 观察面板行数"] == 10 and exhibits["编号图卡数"] == 1
              and exhibits["柱状图轴标签数"] == 2 and exhibits["三重剥离条数"] == 3
              and exhibits["责任链条断点数"] == 13 and exhibits["未采用数字留痕条目"] == 1)
rep["gates"]["G3_design"] = {"detector_p0": det.get("p0"), "detector_findings": findings,
                             "viewports_failing": vp_fail, "viewports_tested": len(vp),
                             "axis_labels": exhibits["柱状图轴标签数"],
                             "pass": det.get("p0") == 0 and not findings and not vp_fail and anchors_ok}
rep["design"] = {"exhibits": exhibits, "exhibit_gate_pass": anchors_ok,
                 "ledger_tokens": {"page": "#ECEFF5", "card": "#FBFCFE",
                                   "accent_cobalt": "#1B3A6B", "gold": "#A9741F"}}

# G4 PPT
pj = json.loads((HERE / "validation_pptx_structure_v3.json").read_text())
s = pj.get("summary", {})
rep["gates"]["G4_ppt"] = {"summary": s, "slides": len(prs.slides._sldIdLst),
                          "notes_on_all": all(x.has_notes_slide and x.notes_slide.notes_text_frame.text.strip()
                                              for x in prs.slides),
                          "pass": s.get("error", 1) == 0 and s.get("warning", 1) == 0}

# G5 PDF + 逐页栅格化留证
import fitz
doc = fitz.open(PDF)
mm = 2.834645
box = fitz.Rect(13 * mm, 14 * mm, (210 - 13) * mm, (297 - 14) * mm)
viol = 0
for pg in doc:
    for b in pg.get_text("blocks"):
        r = fitz.Rect(b[:4])
        if r.y0 > 800: continue
        if r.x0 < box.x0 - 3 or r.x1 > box.x1 + 3 or r.y0 < box.y0 - 3 or r.y1 > box.y1 + 3: viol += 1
FOOT = "98wiki ｜ 智见 / 行业研究报告"
n = doc.page_count
pngs = sorted((HERE / "shots_v3").glob("pdf_v3_p*.png"))
probes = {
  "表1 八段成本树": ["获客","核验","匹配","谈判","带看","按揭","过户","售后"],
  "表2 责任边界矩阵": ["纯 AI 建议（无人工复核）","AI 建议＋持牌经纪人复核","纯人工建议"],
  "表3 三类场景": ["标准化最高、单笔金额小、重复率高","含核验、赎楼、按揭三项持牌环节","继承析产、共有产权、法拍、家庭决策分歧"],
  "表4 Q1—Q5 预言": ["零佣方合同文本中出现按项收费价目","履约主体取得对应执业资质且独立于平台","独家委托合同采用率仍接近零"],
  "表5 观察面板": ["按项收费价目表与验收标准是否公布","人机交接率（AI 可端到端完成环节占比）","单均全链条成本",
    "人均成交套数与服务者净留存率","AI 使用服务者覆盖率","核验纠纷率、赔付金额、返工率",
    "居间责任分类与强制 E&O／营业保证金规则是否出台","AI 建议责任归属的处罚或判例",
    "「AI 居间建议」可保性与承保定价","独家委托合同采用率"],
}
import re as _re
pgtext = [_re.sub(r"\s+", "", doc[i].get_text()) for i in range(n)]
table_single_page = {}
for name, labels in probes.items():
    counts = [sum(1 for l in labels if _re.sub(r"\s+", "", l) in pg) for pg in pgtext]
    table_single_page[name] = {"rows_expected": len(labels), "max_rows_on_one_page": max(counts),
                               "page": counts.index(max(counts)) + 1 if max(counts) else None,
                               "intact": max(counts) == len(labels)}
tables_ok = all(v["intact"] for v in table_single_page.values())
rep["gates"]["G5_pdf"] = {"pages": n, "content_box_violations": viol,
                         "tables_single_page": table_single_page, "tables_intact": tables_ok,
                         "footer_on_body_pages": all(FOOT in doc[i].get_text() for i in range(1, n - 1)),
                         "cover_unfooted": FOOT not in doc[0].get_text(),
                         "back_unfooted": FOOT not in doc[n - 1].get_text(),
                         "rasterized_pages": len(pngs),
                         "pass": viol == 0 and len(pngs) == n and tables_ok
                                 and all(FOOT in doc[i].get_text() for i in range(1, n - 1))}

for f in [MD.name, HTML.name, PDF.name, PPTX.name, "build_deck_v3.py", "gates_v3.py"]:
    p = HERE / f
    rep["artifacts"][f] = {"sha256": sha(p), "bytes": p.stat().st_size}
rep["all_pass"] = all(g["pass"] for g in rep["gates"].values())
(HERE / "gates_report_v3.json").write_text(json.dumps(rep, ensure_ascii=False, indent=1))
print(json.dumps({k: v["pass"] for k, v in rep["gates"].items()}, ensure_ascii=False))
print("new_in_html:", extra, "| new_in_pptx:", extra_p)
print("viewports:", len(vp), "| exhibits:", exhibits)
print("pdf pages:", n, "| rasterized:", len(pngs), "| all_pass:", rep["all_pass"])
