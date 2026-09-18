#!/usr/bin/env python3
"""token_scan.py — 数字纪律三重扫描（对齐 brief.md 数字纪律）。

 A) 量值脱敏禁例 token：全文 0 命中（页面文本 + 表格 + chart XML + notes）
 B) 锚点数字清单：全部在场
 C) 反向扫描：PPTX 中出现的每个阿拉伯数字 token（剔除结构豁免）必须能在 md 中找到
输出 validation/token_scan.json + .md
"""
import json, os, re, sys, datetime
from pptx import Presentation
from pptx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
W = os.path.dirname(os.path.dirname(HERE))
SRC_MD = os.path.join(W, "角色重写_十二主体影响_v5.4.md")
PPTX = os.path.join(HERE, "pptx", "deck_v54.pptx")
OUTD = os.path.join(W, "validation")
os.makedirs(OUTD, exist_ok=True)

md = open(SRC_MD, encoding="utf-8").read()

FORBIDDEN = [
    # 量值脱敏禁例：城市二手成交绝对套数（含不带千分位变体）
    "4,932", "4932", "2,392", "2392", "5,364", "5364", "5,770", "5770",
    "6,039", "6039", "7,920", "7920", "2,606", "2606", "3,097", "3097",
    "4,684", "4684", "3,366", "3366", "2,429", "2429",
]
ANCHORS = [
    "4247", "3614", "+17.5%", "-4bp", "1814", "8.7", "3 万", "100-200",
    "4.5 万", "45-55%", "50%", "30%", "55-60%", "70%",
    "157M", "99.2%", "35-42", "-6.3", "-5.4", "-8.7", "-4.5", "-8.6",
    "36.29", "-4.2%", "-11.8%", "-19.2%", "1.5%", "52.9", "73.5",
    "37.90", "37.56", "37.68", "37.74", "37.44", "37.01", "36.72",
    "3.05%", "3.06%", "-3.4%", "-13.5%", "-2.6%", "-17.2", "-18.0",
    "+1.4%", "+0.6", "+2.0", "-2.3%", "-1.8%", "-4.1", "-11.6%",
    "75-80%", "2-4%", "36 万亿", "1.61", "2027-03", "2026Q4", "2027H1",
    "30%-65%", "35%-60%", "20%-55%", "11 亿", "59 亿", "22,000", "550 亿",
    "1,100", "115 亿", "50%", "40 年", "30 年", "2026-06", "2026-07",
    "3-5 年", "3-5 家", "48 小时", "2027-03-01", "10 亿", "25%", "5%",
    "20%", "15 年", "70%", "80%", "10-15 年", "20 年", "2026-02",
]

# ---- 提取 PPTX 全文本（含表格、chart XML、notes）----
prs = Presentation(PPTX)
texts = []
chart_xmls = []
for i, slide in enumerate(prs.slides, 1):
    def walk(shapes):
        for shp in shapes:
            if shp.shape_type == 6:  # group
                walk(shp.shapes)
                continue
            if getattr(shp, "has_text_frame", False):
                texts.append((i, "shape", shp.text_frame.text))
            if getattr(shp, "has_table", False):
                for r in shp.table.rows:
                    for c in r.cells:
                        texts.append((i, "table", c.text))
            if shp.has_chart:
                chart_xmls.append((i, shp.chart._chartSpace.xml))
    walk(slide.shapes)
    if slide.has_notes_slide:
        texts.append((i, "notes", slide.notes_slide.notes_text_frame.text))

all_text = "\n".join(t for _, _, t in texts)
chart_all = "\n".join(x for _, x in chart_xmls)

# ---- A) 禁例扫描 ----
forb_hits = []
for tok in FORBIDDEN:
    if tok in all_text:
        pages = sorted({i for i, _, t in texts if tok in t})
        forb_hits.append((tok, pages))
    if tok in chart_all:
        pages = sorted({i for i, x in chart_xmls if tok in x})
        forb_hits.append((tok + " [chart]", pages))

# ---- B) 锚点在场 ----
missing_anchors = [a for a in ANCHORS if a not in all_text and a not in chart_all]

# ---- C) 反向扫描 ----
STRUCT_EXEMPT = re.compile(r"^(0[1-9]|[12][0-9]|3[0-3])$")  # 页码 NN / 33
def exempt(tok, page):
    if STRUCT_EXEMPT.match(tok):
        return True
    if tok in ("①②③④⑤⑥⑦⑧⑨⑩⑪⑫⑬",):
        return True
    return False

num_re = re.compile(r"\d[\d,\.%]*")
unknown = {}
for i, kind, t in texts:
    for m in num_re.finditer(t):
        tok = m.group(0).rstrip("%")
        if exempt(tok, i):
            continue
        probe = tok.replace(",", "")
        if probe not in md and tok not in md:
            unknown.setdefault(tok, []).append(f"S{i}:{kind}")

chart_num_re = re.compile(r">(\d[\d,\.]*)<")
for i, x in chart_xmls:
    for m in chart_num_re.finditer(x):
        tok = m.group(1)
        if tok in md:
            continue
        unknown.setdefault(tok + " (chart)", []).append(f"S{i}:chart")

rep = {
    "generated_at": datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
    "forbidden_scan": {"tokens": len(FORBIDDEN), "hits": forb_hits,
                       "verdict": "PASS (0 hits)" if not forb_hits else "FAIL"},
    "anchor_scan": {"checked": len(ANCHORS), "missing": missing_anchors,
                    "verdict": "PASS (all present)" if not missing_anchors else "FAIL"},
    "reverse_scan": {"unknown_tokens": {k: v[:6] for k, v in sorted(unknown.items())},
                     "count": len(unknown),
                     "verdict": "PASS" if not unknown else "REVIEW"},
}
with open(os.path.join(OUTD, "token_scan.json"), "w", encoding="utf-8") as f:
    json.dump(rep, f, ensure_ascii=False, indent=2)
md_lines = [
    "# token 纪律扫描", "",
    f"- A 禁例（量值脱敏）：{len(FORBIDDEN)} 个 token，命中 {len(forb_hits)} → **{rep['forbidden_scan']['verdict']}**",
    f"- B 锚点清单：{len(ANCHORS)} 项，缺失 {len(missing_anchors)} → **{rep['anchor_scan']['verdict']}**",
    f"- C 反向扫描：未溯源数字 token {len(unknown)} 个 → **{rep['reverse_scan']['verdict']}**",
]
for k, v in sorted(unknown.items()):
    md_lines.append(f"  - `{k}` @ {', '.join(v[:6])}")
for a in missing_anchors:
    md_lines.append(f"  - MISSING ANCHOR: {a}")
with open(os.path.join(OUTD, "token_scan.md"), "w", encoding="utf-8") as f:
    f.write("\n".join(md_lines) + "\n")
print(f"forbidden hits={len(forb_hits)} | missing anchors={len(missing_anchors)} {missing_anchors[:8]}"
      f" | unknown tokens={len(unknown)}")
for k, v in sorted(unknown.items())[:30]:
    print(f"  ? {k} @ {v[:4]}")
